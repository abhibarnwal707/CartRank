import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from cartrank.eval import demo_labels


def _rows():
    return [
        [1, "stapler", "A1", "Desk Stapler", "Acme", 9.5, "Office > Staplers", "metal", ""],
        [1, "stapler", "A2", "Paper Clips", "Acme", 2.0, "Office > Clips", "steel", ""],
        [2, "pens", "A3", "Gel Pens", "Bic", 4.0, "Office > Pens", "blue", ""],
    ]


@pytest.fixture
def sheet(tmp_path, monkeypatch):
    monkeypatch.setattr(demo_labels, "SHEET", tmp_path / "sheet.xlsx")
    monkeypatch.setattr(demo_labels, "RUNS", tmp_path / "runs.json")
    monkeypatch.setattr(demo_labels, "RESULTS", tmp_path)
    demo_labels.write_sheet(_rows())
    return tmp_path


def test_sheet_layout_and_blind(sheet):
    wb = load_workbook(sheet / "sheet.xlsx")
    assert wb.sheetnames == ["Instructions", "Label"]
    df = pd.read_excel(sheet / "sheet.xlsx", sheet_name="Label")
    assert list(df.columns) == demo_labels.HEADERS
    assert len(df) == 3 and df.label.isna().all()
    assert not {"system", "bm25", "final"} & set(map(str.lower, df.columns))
    assert wb["Label"].data_validations.dataValidation[0].formula1 == '"1,0"'


def test_score_rejects_unlabelled_and_computes_precision(sheet):
    (sheet / "runs.json").write_text(
        json.dumps(
            {
                "meta": {"tag": "ft", "hybrid_k": 10, "hybrid_depth": 200},
                "queries": {
                    "1": {"query": "stapler", "bm25": ["A1", "A2"], "final": ["A1"]},
                    "2": {"query": "pens", "bm25": ["A3"], "final": ["A3"]},
                },
            }
        )
    )
    with pytest.raises(SystemExit):
        demo_labels.score()
    wb = load_workbook(sheet / "sheet.xlsx")
    for r, v in zip((2, 3, 4), (1, 0, 1), strict=True):
        wb["Label"].cell(row=r, column=11, value=v)
    wb.save(sheet / "sheet.xlsx")
    demo_labels.score()
    out = (sheet / "demo_transfer.md").read_text()
    assert "| BM25 | 0.750 | - |" in out  # q1: (1+0)/2 = 0.5, q2: 1.0 -> mean 0.75
    assert "| Final | 1.000 |" in out
