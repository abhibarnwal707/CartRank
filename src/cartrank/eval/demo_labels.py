"""Phase 3 transfer check: build a blind labelling sheet, then score it.

python -m cartrank.eval.demo_labels make <tag> <encoder_path>   -> xlsx for the owner to label
python -m cartrank.eval.demo_labels score                       -> results/demo_transfer.md

The sheet pools each query's BM25 top-10 and final-system top-10, de-duplicated and shuffled,
with no system names, so the labeller cannot tell which system produced a product. System
identity lives only in results/demo_runs.json.
"""

import json
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from cartrank.data.paths import RESULTS, ROOT

QUERIES = ROOT / "evals" / "queries" / "demo_queries.jsonl"
RUNS = RESULTS / "demo_runs.json"
SHEET = ROOT / "evals" / "queries" / "demo_labelling_sheet.xlsx"
FONT = "Arial"
HEADERS = [
    "row",
    "query_id",
    "query",
    "asin",
    "title",
    "brand",
    "price_usd",
    "category",
    "about_this_item",
    "image_url",
    "label",
    "note",
]
WIDTHS = [6, 9, 28, 13, 60, 16, 10, 34, 70, 30, 9, 28]


def make(tag: str, model_path: str, top: int = 10, seed: int = 0) -> None:
    from cartrank.search.catalog_index import CatalogSearch, hybrid_setting

    k, depth = hybrid_setting(tag)
    cs = CatalogSearch(tag, model_path, k, depth)
    qs = [json.loads(line) for line in QUERIES.read_text().splitlines()]
    runs = {"meta": {"tag": tag, "hybrid_k": k, "hybrid_depth": depth, "top": top}, "queries": {}}
    for q in qs:
        runs["queries"][str(q["id"])] = {
            "query": q["query"],
            "bm25": cs.bm25.search(q["query"], top),
            "final": cs.final(q["query"])[:top],
        }
    RUNS.write_text(json.dumps(runs, indent=1))

    rng = np.random.default_rng(seed)
    cat = cs.cat.set_index("parent_asin")
    rows = []
    for q in qs:
        r = runs["queries"][str(q["id"])]
        pool = list(dict.fromkeys(r["bm25"] + r["final"]))
        for asin in rng.permutation(pool):
            c = cat.loc[asin]
            rows.append(
                [
                    q["id"],
                    q["query"],
                    asin,
                    c.title,
                    c.store or "",
                    float(c.price),
                    " > ".join(list(c.categories)[1:]),
                    " ".join(list(c.features))[:300],
                    c.image_url or "",
                ]
            )
    write_sheet(rows)
    print(f"{len(rows)} rows to label across {len(qs)} queries -> {SHEET}")


def write_sheet(rows: list[list]) -> None:
    wb = Workbook()
    guide = wb.active
    guide.title = "Instructions"
    lab = wb.create_sheet("Label")
    n = len(rows)
    last = n + 1

    lines = [
        ("CartRank demo labelling sheet", True),
        ("", False),
        (
            "Goal: judge whether each product is a good result for the search query next to it.",
            False,
        ),
        (
            "Work on the 'Label' sheet. Fill ONLY the yellow 'label' column (and 'note' if you like).",
            False,
        ),
        ("", False),
        ("Label values (use the dropdown):", True),
        (
            "   1 = relevant: a shopper who typed this query would be happy to see this product.",
            False,
        ),
        (
            "   0 = not relevant: wrong product type, wrong key attribute, or violates a stated constraint.",
            False,
        ),
        ("", False),
        ("Tips:", True),
        (
            "   - Judge each row against its own query. The same product can be 1 for one query and 0 for another.",
            False,
        ),
        ("   - If a query says 'without X' or 'not Y', a product with X or Y is 0.", False),
        ("   - Typos: judge the intended meaning ('stappler' means stapler).", False),
        (
            "   - Borderline (right type but a different size, colour or brand than asked)? Use your own",
            False,
        ),
        ("     consistent rule and write it in 'note' once, then apply it everywhere.", False),
        (
            "   - Do not guess which search engine returned a product; the order is shuffled and unlabelled.",
            False,
        ),
        (
            "   - Use the title, brand, price and 'about_this_item'; open image_url only if unsure.",
            False,
        ),
        (
            "   - Rows highlighted yellow are still unlabelled. Save as .xlsx (keep the file name).",
            False,
        ),
        ("", False),
        ("Example row (illustrative, not part of the data):", True),
        (
            "   query 'blue gel pens' | 'Pilot G2 Gel Pens, Blue, 12 pack' -> 1  |  'Blue Notebook, wide ruled' -> 0",
            False,
        ),
        ("", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = guide.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=14 if i == 1 else 11)
    r0 = len(lines) + 1
    guide.cell(row=r0, column=1, value="Progress").font = Font(name=FONT, bold=True)
    guide.cell(row=r0 + 1, column=1, value="Rows labelled:").font = Font(name=FONT)
    guide.cell(row=r0 + 1, column=2, value=f"=COUNTA(Label!K2:K{last})").font = Font(name=FONT)
    guide.cell(row=r0 + 2, column=1, value="Rows total:").font = Font(name=FONT)
    guide.cell(row=r0 + 2, column=2, value=f"=COUNTA(Label!A2:A{last})").font = Font(name=FONT)
    guide.cell(row=r0 + 3, column=1, value="Percent done:").font = Font(name=FONT)
    pct = guide.cell(row=r0 + 3, column=2, value=f"=IF(B{r0 + 2}=0,0,B{r0 + 1}/B{r0 + 2})")
    pct.number_format = "0.0%"
    pct.font = Font(name=FONT)
    guide.column_dimensions["A"].width = 110
    guide.column_dimensions["B"].width = 12

    for j, h in enumerate(HEADERS, 1):
        c = lab.cell(row=1, column=j, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for i, r in enumerate(rows, 1):
        for j, v in enumerate([i, *r], 1):
            c = lab.cell(row=i + 1, column=j, value=v)
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=j in (3, 5, 8, 9), vertical="top")
        lab.cell(row=i + 1, column=7).number_format = "$#,##0.00"
    for j, w in enumerate(WIDTHS, 1):
        lab.column_dimensions[get_column_letter(j)].width = w
    lab.freeze_panes = "E2"
    lab.auto_filter.ref = f"A1:L{last}"
    dv = DataValidation(
        type="list",
        formula1='"1,0"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Label",
        error="Enter 1 (relevant) or 0 (not relevant).",
    )
    lab.add_data_validation(dv)
    dv.add(f"K2:K{last}")
    lab.conditional_formatting.add(
        f"K2:K{last}",
        FormulaRule(formula=['K2=""'], fill=PatternFill("solid", bgColor="FFFF00")),
    )
    SHEET.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SHEET)


def score() -> None:
    df = pd.read_excel(SHEET, sheet_name="Label")
    if df.label.isna().any():
        raise SystemExit(f"{int(df.label.isna().sum())} rows are still unlabelled")
    if not set(df.label.astype(int)) <= {0, 1}:
        raise SystemExit("labels must be 0 or 1")
    lab = {(int(q), a): int(v) for q, a, v in zip(df.query_id, df.asin, df.label, strict=True)}
    runs = json.loads(RUNS.read_text())
    qtype = {}
    for line in QUERIES.read_text().splitlines():
        q = json.loads(line)
        qtype[q["id"]] = q["type"]
    rows = []
    for qid, r in runs["queries"].items():
        rows.append(
            {
                "type": qtype[int(qid)],
                "bm25": np.mean([lab[(int(qid), a)] for a in r["bm25"]]),
                "final": np.mean([lab[(int(qid), a)] for a in r["final"]]),
            }
        )
    d = pd.DataFrame(rows)
    d["diff"] = d.final - d.bm25
    rng = np.random.default_rng(0)
    boot = [rng.choice(d["diff"].to_numpy(), len(d)).mean() for _ in range(10_000)]
    lo, hi = np.percentile(boot, [2.5, 97.5])
    by_type = d.groupby("type").agg(
        n=("diff", "size"), bm25=("bm25", "mean"), final=("final", "mean")
    )
    meta = runs["meta"]
    out = (
        "# Demo catalog transfer check (Office_Products, 50 queries)\n\n"
        "Precision@10 on a small sanity check, **not** a headline metric. The ESCI-trained models "
        "are applied to a catalog with no relevance labels; the owner hand-labelled the top-10 of "
        "each system (blind: products were pooled, shuffled and unlabelled by system).\n\n"
        f"Final system: hybrid (BM25 + `{meta['tag']}` dense, RRF k={meta['hybrid_k']}, "
        f"depth={meta['hybrid_depth']}) + cross-encoder rerank of the top 50.\n\n"
        "| System | Precision@10 | Δ vs BM25 |\n|---|---|---|\n"
        f"| BM25 | {d.bm25.mean():.3f} | - |\n"
        f"| Final | {d.final.mean():.3f} | {d['diff'].mean():+.3f} |\n\n"
        f"Paired bootstrap over the 50 queries (10,000 resamples, seed 0): the mean difference "
        f"has a 95% interval of [{lo:+.3f}, {hi:+.3f}]. Per query, the final system is better on "
        f"{int((d['diff'] > 0).sum())}, tied on {int((d['diff'] == 0).sum())} and worse on "
        f"{int((d['diff'] < 0).sum())}.\n\n"
        "## By query type\n\n"
        "| Type | Queries | BM25 P@10 | Final P@10 |\n|---|---|---|---|\n"
        + "".join(
            f"| {t} | {int(r.n)} | {r.bm25:.3f} | {r.final:.3f} |\n" for t, r in by_type.iterrows()
        )
        + f"\nQueries: {len(d)}; labelled products: {len(df)}. Binary labels (1 relevant, 0 not). "
        "Per-type groups have 4 to 15 queries, so treat them as indicative only.\n"
    )
    (RESULTS / "demo_transfer.md").write_text(out)
    print(out)


if __name__ == "__main__":
    if sys.argv[1] == "make":
        make(sys.argv[2], sys.argv[3])
    else:
        score()
